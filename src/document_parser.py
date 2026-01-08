from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional
import logging

logger = logging.getLogger("policy_parser")


@dataclass
class Document:
    content: str
    metadata: Dict[str, str]
    source: str
    page_numbers: Optional[List[int]] = None


class DocumentParser:
    def __init__(self, header_footer_ratio: float = 0.08):
        self.supported_formats = {".pdf", ".docx", ".txt"}
        self.header_footer_ratio = header_footer_ratio

    def parse(self, file_path: Path | str) -> Document:
        file_path = Path(file_path)
        suffix = file_path.suffix.lower()

        if suffix not in self.supported_formats:
            raise ValueError(f"Unsupported format: {suffix}")

        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        if suffix == ".docx":
            return self._parse_docx(file_path)
        return self._parse_txt(file_path)

    def _parse_pdf(self, file_path: Path) -> Document:
        import fitz  # PyMuPDF

        pdf = fitz.open(str(file_path))
        page_count = pdf.page_count

        content_parts: List[str] = []
        page_numbers: List[int] = []

        for page_idx in range(page_count):
            page = pdf[page_idx]
            page_num = page_idx + 1

            blocks = page.get_text("blocks")
            if not blocks:
                continue

            page_h = page.rect.height
            top_cut = page_h * self.header_footer_ratio
            bot_cut = page_h * (1.0 - self.header_footer_ratio)

            filtered = []
            for b in blocks:
                x0, y0, x1, y1, text = b[0], b[1], b[2], b[3], b[4]
                if not text or not text.strip():
                    continue
                if y1 < top_cut or y0 > bot_cut:
                    continue
                filtered.append((x0, y0, x1, y1, text.strip()))

            if not filtered:
                continue

            filtered.sort(key=lambda t: (t[1], t[0]))
            page_text = "\n".join(t[4] for t in filtered)

            content_parts.append(f"[Page {page_num}]\n{page_text}")
            page_numbers.append(page_num)

        metadata = {
            "title": (pdf.metadata or {}).get("title", "") or "",
            "author": (pdf.metadata or {}).get("author", "") or "",
            "pages": str(page_count),
            "format": "pdf",
            "parser": "pymupdf",
        }

        pdf.close()
        logger.info(f"Parsed PDF: {file_path.name}, pages={page_count}, kept_pages={len(page_numbers)}")

        return Document(
            content="\n\n".join(content_parts),
            metadata=metadata,
            source=str(file_path),
            page_numbers=page_numbers,
        )

    def _parse_docx(self, file_path: Path) -> Document:
        from docx import Document as DocxDocument

        docx = DocxDocument(str(file_path))
        parts: List[str] = []

        for p in docx.paragraphs:
            text = (p.text or "").strip()
            if not text:
                continue

            style = (p.style.name or "").lower()
            if style.startswith("heading"):
                parts.append(f"\n## {text}\n")
            else:
                parts.append(text)

        for table in docx.tables:
            parts.append("\n" + self._parse_table(table) + "\n")

        metadata = {
            "title": file_path.stem,
            "paragraphs": str(len(docx.paragraphs)),
            "tables": str(len(docx.tables)),
            "format": "docx",
            "parser": "python-docx",
        }

        logger.info(f"Parsed DOCX: {file_path.name}, paragraphs={len(docx.paragraphs)}, tables={len(docx.tables)}")

        return Document(content="\n".join(parts), metadata=metadata, source=str(file_path))

    def _parse_table(self, table) -> str:
        rows = []
        for row in table.rows:
            cells = [" ".join((cell.text or "").split()) for cell in row.cells]
            rows.append(" | ".join(cells).strip())
        return "\n".join(r for r in rows if r)

    def _parse_txt(self, file_path: Path) -> Document:
        try:
            content = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = file_path.read_text(encoding="utf-8", errors="ignore")

        metadata = {"title": file_path.stem, "format": "txt", "parser": "plaintext"}
        logger.info(f"Parsed TXT: {file_path.name}, chars={len(content)}")

        return Document(content=content, metadata=metadata, source=str(file_path))
